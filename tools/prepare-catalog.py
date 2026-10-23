from pathlib import Path
import json,hashlib
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]
workspace=root.parent
book=workspace/'延河课堂-全部课程信息.xlsx'
modern=workspace/'延河课堂全部课程/courses-all.json'
courses={}
if book.exists():
 w=load_workbook(book,read_only=True,data_only=True)
 it=w['全部课程'].iter_rows(values_only=True);headers=next(it)
 for values in it:
  r=dict(zip(headers,values));id=str(r.get('课程ID') or '').strip()
  if not id.isdigit():continue
  item={'id':id,'platform':'legacy','title':r.get('课程名称') or '', 'teacher':r.get('授课教师') or '', 'term':r.get('学年学期') or '', 'college':r.get('开课学院') or '', 'code':r.get('课程代码') or ''}
  courses['legacy:'+id]=item
 w.close()
if modern.exists():
 for r in json.loads(modern.read_text(encoding='utf8'))['courses']:
  id=str(r['id']);courses['modern:'+id]={k:r.get(k,'') for k in ['id','title','teacher','term','college']}|{'platform':'modern'}
out=root/'local-data';out.mkdir(exist_ok=True)
data={'schema':1,'courses':list(courses.values()),'source_files':[p.name for p in [book,modern] if p.exists()]}
file=out/'catalog.json';file.write_text(json.dumps(data,ensure_ascii=False,separators=(',',':')),encoding='utf8')
print(json.dumps({'records':len(courses),'legacy':sum(r['platform']=='legacy' for r in courses.values()),'modern':sum(r['platform']=='modern' for r in courses.values()),'bytes':file.stat().st_size,'sha256':hashlib.sha256(file.read_bytes()).hexdigest()},ensure_ascii=False))
